#!/usr/bin/env python3
"""IT 대시보드 엑셀 업로드 처리 — GitHub Actions(.github/workflows/it-upload.yml)가 호출한다.

담당자가 GitHub 웹에서 엑셀을 올리면 이 스크립트가 검증하고 화면용 JSON을 만든다.
원본(PTT Raw·CMS·Readiness)은 다루지 않는다 — 결과물 엑셀만 받는다.

    python3 it/_tools/it_upload.py export          # data/npi-product-status.json → _upload/IT_product_status.xlsx
    python3 it/_tools/it_upload.py import          # _upload/IT_product_status.xlsx → data/npi-product-status.json
    python3 it/_tools/it_upload.py gr <prev.xlsx> --prev-date YYYYMMDD   # Global Request.xlsx 검증 + data/gr-changes.json

검증에 실패하면 종료코드 1과 함께 사유를 출력하고 아무것도 쓰지 않는다.
"""
from __future__ import annotations

import argparse
import json
import os
import subprocess
import sys
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

IT = Path(__file__).resolve().parent.parent
STATUS_JSON = IT / 'data' / 'npi-product-status.json'
STATUS_XLSX = IT / '_upload' / 'IT_product_status.xlsx'
GR_XLSX = IT / 'Global Request.xlsx'
GR_CHANGES = IT / 'data' / 'gr-changes.json'
GR_TITLES = IT / 'data' / 'gr-titles.json'
KST = timezone(timedelta(hours=9))

SHEET = 'IT 제품 현황'
INFO_SHEET = '정보'
RC_PREFIX = 'RC: '
RC_FIELDS = ['Channel DAM Workflow', 'PDR Use', 'Spec Status', 'UFN', 'Key Feature',
             'Spec Assignment', 'Local Asset Review', 'Live URL B2C', 'Live URL SMB',
             'Live URL D2B2C', 'Live URL Partners', 'Live URL Education']
# (JSON 키, 엑셀 머리글, 설명) — 순서가 곧 엑셀 열 순서. 2행에 키를 적어 머리글을 바꿔도 읽힌다.
COLUMNS = [
    ('_group', '그룹', '화면 그룹명 (예: NPI 7월). 같은 그룹명·색상끼리 묶인다'),
    ('_color', '그룹 색상', 'green / orange / red / blue / gray'),
    ('category', '구분', 'NPI / MOD'),
    ('targetLaunchDate', '출시 예정일', 'YYYY-MM-DD'),
    ('region', 'Region', ''),
    ('sub', '법인', ''),
    ('locale', 'Locale', '예: IT-it'),
    ('model', '모델', ''),
    ('salesModelKey', 'Sales Model Key', '비우면 Locale-모델로 채움'),
    ('stage', '단계', 'Live / In Progress / Client Review / Clarify / Precheck / Cancelled / 기타'),
    ('stageForce', '단계 고정', '비우면 자동. live면 라이브 확정'),
    ('stageFallback', '단계 기본값', ''),
    ('pttId', 'PTT ID', ''),
    ('pttStatus', 'PTT 상태', ''),
    ('targetStagingDate', 'STG 목표일', ''),
    ('expectedLocalTargetDate', '법인 예상일', ''),
    ('hasStg', 'STG 있음', 'Y / N'),
    ('stgUrl', 'STG URL', ''),
    ('liveUrl', '라이브 URL', 'https:// 로 시작'),
    ('cmsLive', 'CMS', '라이브 / 등록·미게시 / PIM 미등록'),
    ('pdpStatus', 'PDP 이슈', ''),
    ('pimChecked', 'PIM 확인', 'Y / N'),
    ('readinessCheck', 'Readiness', 'Ready / Not Ready'),
    ('readinessTip', 'Readiness 메모', ''),
    ('detailKr', '상세(KR)', ''),
    ('detailEn', '상세(EN)', ''),
    ('statusTip', '상태 메모', ''),
] + [(RC_PREFIX + f, RC_PREFIX + f, 'Readiness 원본 필드') for f in RC_FIELDS]
BOOL_KEYS = {'hasStg', 'pimChecked'}
REQUIRED = {'_group', 'locale', 'model'}
COLORS = {'green', 'orange', 'red', 'blue', 'gray', 'grey', 'yellow', 'purple'}
INFO_KEYS = [('sourceFile', '기준 목록 파일'), ('pttRawFile', 'PTT Raw 파일'),
             ('cmsFile', 'CMS 파일'), ('readinessFile', 'Readiness 파일')]


class Invalid(Exception):
    pass


def _str(v) -> str:
    if v is None:
        return ''
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _summary(text: str) -> None:
    print(text)
    path = os.environ.get('GITHUB_STEP_SUMMARY')
    if path:
        with open(path, 'a', encoding='utf-8') as f:
            f.write(text + '\n')


# ── 제품 현황: JSON → 엑셀 ─────────────────────────────────────────────
def export_status(src: Path = STATUS_JSON, out: Path = STATUS_XLSX) -> None:
    d = json.loads(src.read_text(encoding='utf-8'))
    labels = {s['key']: s['label'] for s in d.get('stages') or []}
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append([c[1] for c in COLUMNS])
    ws.append([c[0] for c in COLUMNS])
    for g in d['groups']:
        for r in g['rows']:
            rf = r.get('readinessFields') or {}
            vals = []
            for key, _, _ in COLUMNS:
                if key == '_group':
                    v = g.get('status', '')
                elif key == '_color':
                    v = g.get('colorClass', '')
                elif key.startswith(RC_PREFIX):
                    v = rf.get(key[len(RC_PREFIX):], '')
                elif key == 'stage':
                    v = labels.get(r.get('stage') or 'etc', r.get('stage') or '')
                elif key in BOOL_KEYS:
                    v = 'Y' if r.get(key) else 'N'
                else:
                    v = r.get(key)
                    v = '' if v is None else v
                vals.append(v)
            ws.append(vals)
    head_fill = PatternFill('solid', fgColor='A50034')
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = head_fill
        c.alignment = Alignment(vertical='center', wrap_text=True)
    for c in ws[2]:
        c.font = Font(italic=True, color='999999', size=8)
    for i, (key, label, note) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = 34 if key in ('detailKr', 'detailEn', 'liveUrl', 'stgUrl', 'readinessTip') else max(10, min(24, len(label) * 2 + 2))
    ws.freeze_panes = 'C3'
    ws.auto_filter.ref = f'A2:{get_column_letter(len(COLUMNS))}{ws.max_row}'

    info = wb.create_sheet(INFO_SHEET)
    info.append(['항목', '값', '설명'])
    info.append(['기준 시각', d.get('generatedAt', ''), '업로드하면 자동으로 바뀜'])
    for key, label in INFO_KEYS:
        info.append([label, d.get(key, ''), key])
    info.append([])
    info.append(['열', '키', '입력 규칙'])
    for key, label, note in COLUMNS:
        info.append([label, key, note])
    info.column_dimensions['A'].width = 22
    info.column_dimensions['B'].width = 46
    info.column_dimensions['C'].width = 60
    for row in (1, len(INFO_KEYS) + 4):
        for c in info[row]:
            c.font = Font(bold=True)
    if out.exists() and _cell_values(out) == _cell_values(wb):
        print(f'양식 변경 없음: {out.name}')   # 다시 저장하면 내부 시각만 바뀌어 빈 커밋이 생긴다
        return
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f'양식 기록: {out.name} ({ws.max_row - 2}행)')


def _cell_values(src) -> list:
    wb = src if isinstance(src, Workbook) else load_workbook(src)
    # 메모리의 ''는 저장 후 None으로 읽히므로 같게 본다. 끝의 빈 행·열도 무시한다.
    out = []
    for ws in wb.worksheets:
        rows = [[None if v == '' else v for v in r] for r in ws.iter_rows(values_only=True)]
        rows = [r[:max((i + 1 for i, v in enumerate(r) if v is not None), default=0)] for r in rows]
        while rows and not rows[-1]:
            rows.pop()
        out.append((ws.title, rows))
    return out


# ── 제품 현황: 엑셀 → JSON ─────────────────────────────────────────────
def read_status_xlsx(path: Path, prev: dict) -> dict:
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise Invalid(f'엑셀 파일을 열 수 없습니다: {e}')
    if SHEET not in wb.sheetnames:
        raise Invalid(f"'{SHEET}' 시트가 없습니다. 시트 이름을 바꾸지 마세요 (현재: {wb.sheetnames})")
    ws = wb[SHEET]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 3:
        raise Invalid('데이터 행이 없습니다 (1행 머리글, 2행 키, 3행부터 데이터).')
    keys = [_str(c) for c in rows[1]]
    known = {c[0] for c in COLUMNS}
    missing = sorted(REQUIRED - set(keys))
    if missing:
        raise Invalid(f'2행(키 행)에 필수 열이 없습니다: {missing}. 2행은 지우거나 고치지 마세요.')
    unknown = [k for k in keys if k and k not in known]
    if unknown:
        raise Invalid(f'알 수 없는 열 키: {unknown}. 2행 키는 양식 그대로 두세요.')

    stages = prev.get('stages') or []
    by_label = {s['label'].lower(): s['key'] for s in stages}
    by_key = {s['key'].lower(): s['key'] for s in stages}
    errors: list[str] = []
    groups: list[dict] = []
    group_idx: dict[tuple[str, str], dict] = {}
    for n, raw in enumerate(rows[2:], start=3):
        cells = {k: _str(v) for k, v in zip(keys, raw) if k}
        if not any(cells.values()):
            continue
        for k in REQUIRED:
            if not cells.get(k):
                errors.append(f'{n}행: {dict(COLUMNS_LABEL)[k]} 값이 비어 있습니다')
        stage_in = cells.get('stage', '')
        stage = by_label.get(stage_in.lower()) or by_key.get(stage_in.lower()) or ('etc' if not stage_in else None)
        if stage is None:
            errors.append(f"{n}행: 단계 '{stage_in}'를 알 수 없습니다 (허용: {', '.join(s['label'] for s in stages)})")
        color = cells.get('_color', '').lower() or 'gray'
        if color not in COLORS:
            errors.append(f"{n}행: 그룹 색상 '{cells.get('_color')}'는 쓸 수 없습니다 (허용: green/orange/red/blue/gray)")
        for k in ('liveUrl', 'stgUrl'):
            v = cells.get(k, '')
            if v and not v.startswith(('http://', 'https://')):
                errors.append(f"{n}행: {dict(COLUMNS_LABEL)[k]}는 http(s):// 로 시작해야 합니다 ('{v[:40]}')")
        for k in BOOL_KEYS:
            v = cells.get(k, '').upper()
            if v and v not in ('Y', 'N', 'O', 'X', 'TRUE', 'FALSE', '1', '0'):
                errors.append(f"{n}행: {dict(COLUMNS_LABEL)[k]}는 Y/N 으로 적어 주세요 ('{cells.get(k)}')")

        row: dict = {}
        rf: dict = {}
        for key, _, _ in COLUMNS:
            if key in ('_group', '_color') or key not in cells:
                continue
            v = cells[key]
            if key.startswith(RC_PREFIX):
                if v:
                    rf[key[len(RC_PREFIX):]] = v
            elif key in BOOL_KEYS:
                row[key] = v.upper() in ('Y', 'O', 'TRUE', '1')
            elif key == 'stage':
                row[key] = stage or 'etc'
            else:
                row[key] = v
        if rf:
            row['readinessFields'] = {f: rf.get(f, '') for f in RC_FIELDS}
        if not row.get('salesModelKey'):
            row['salesModelKey'] = f"{row.get('locale', '')}-{row.get('model', '')}"
        gkey = (cells['_group'], color)
        g = group_idx.get(gkey)
        if g is None:
            g = {'status': cells['_group'], 'colorClass': color, 'rows': []}
            group_idx[gkey] = g
            groups.append(g)
        g['rows'].append(row)
    if errors:
        raise Invalid('\n'.join(errors[:50]) + (f'\n… 외 {len(errors) - 50}건' if len(errors) > 50 else ''))
    total = sum(len(g['rows']) for g in groups)
    if total == 0:
        raise Invalid('데이터 행이 0건입니다.')

    out = dict(prev)
    out['groups'] = groups
    out['totalRows'] = total
    out['generatedAt'] = datetime.now(KST).strftime('%Y-%m-%d %H:%M')
    if INFO_SHEET in wb.sheetnames:
        labels = {label: key for key, label in INFO_KEYS}
        for r in wb[INFO_SHEET].iter_rows(values_only=True):
            if r and _str(r[0]) in labels:
                out[labels[_str(r[0])]] = _str(r[1] if len(r) > 1 else '')
    return out


COLUMNS_LABEL = [(k, label) for k, label, _ in COLUMNS]


def _norm(v):
    """빈 값('', None, False, 빈 dict)과 키 없음을 같게 본다 — 엑셀 왕복에서 구분되지 않는 차이."""
    if isinstance(v, dict):
        return {k: _norm(x) for k, x in v.items() if _norm(x) not in ('', {})}
    if isinstance(v, list):
        return [_norm(x) for x in v]
    return '' if v in (None, '', False) else v


def import_status(src: Path = STATUS_XLSX, out: Path = STATUS_JSON) -> int:
    prev = json.loads(out.read_text(encoding='utf-8'))
    try:
        new = read_status_xlsx(src, prev)
    except Invalid as e:
        _summary(f'## ❌ IT 제품 현황 업로드 실패\n\n반영하지 않았습니다. 엑셀을 고쳐 다시 올려 주세요.\n\n```\n{e}\n```')
        return 1
    meta_keys = [k for k, _ in INFO_KEYS]
    if _norm(new['groups']) == _norm(prev['groups']) and all(new.get(k) == prev.get(k) for k in meta_keys):
        _summary('## ✅ IT 제품 현황 — 변경 없음\n\n올린 엑셀이 현재 화면과 같아 그대로 두었습니다.')
        return 0
    before = {r.get('salesModelKey') for g in prev['groups'] for r in g['rows']}
    after = {r.get('salesModelKey') for g in new['groups'] for r in g['rows']}
    out.write_text(json.dumps(new, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    _summary(f"## ✅ IT 제품 현황 반영\n\n- 행 수: {sum(len(g['rows']) for g in prev['groups'])} → {new['totalRows']}\n"
             f"- 추가 {len(after - before)} · 삭제 {len(before - after)}\n- 그룹 {len(new['groups'])}개\n- 기준 시각 {new['generatedAt']}")
    return 0


# ── GR 시트: 검증 + 금주 변경(gr-changes.json) ─────────────────────────
# 아래 read_gr_items·diff_gr_items는 it-dashboard 저장소 scripts/rebuild_gr_changes.py,
# build-reference-workbook.py의 같은 이름 함수를 옮긴 것이다. 규칙을 바꾸면 양쪽을 함께 고친다.
def read_gr_items(xlsx: Path, title_to_task: dict) -> list[dict]:
    wb = load_workbook(xlsx, data_only=True)
    items: list[dict] = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        b1 = ws['B1'].value
        task = title_to_task.get(str(b1).strip() if b1 is not None else '')
        if not task:
            task = name.split(' - ', 1)[-1].strip()
        head = hdr = None
        for i, r in enumerate(rows):
            if r and r[0] == 'Head':
                head = i
                hdr = [str(c).strip() if c is not None else '' for c in r]
                break
        if head is None:
            continue
        n = len(hdr)
        ci = hdr.index('Country') if 'Country' in hdr else 1
        si = hdr.index('Status') if 'Status' in hdr else None
        ui = hdr.index('URL') if 'URL' in hdr else None
        for r in rows[head + 1:]:
            if not r or all(c is None or str(c).strip() == '' for c in r):
                continue
            r = tuple(r) + (None,) * (n - len(r))
            country = str(r[ci] or '').strip()
            if not country:
                continue
            items.append({'task': task, 'countryRaw': country,
                          'status': str(r[si] or '').strip() if si is not None else '',
                          'url': str(r[ui] or '').strip() if ui is not None else ''})
    return items


def normalize_countries(raws: list[str]) -> dict:
    """js-lib/gr-meta.js의 grNormalizeCountryCode를 그대로 호출한다(화면과 같은 규칙)."""
    script = ("const m=require(process.argv[1]);let s='';process.stdin.on('data',d=>s+=d).on('end',()=>{"
              "const o={};for(const v of JSON.parse(s))o[v]=m.grNormalizeCountryCode(v);process.stdout.write(JSON.stringify(o));});")
    out = subprocess.run(['node', '-e', script, str(IT / 'js-lib' / 'gr-meta.js')],
                         input=json.dumps(sorted(set(raws)), ensure_ascii=False),
                         capture_output=True, text=True, check=True).stdout
    return json.loads(out)


def diff_gr_items(prev_items: list[dict], curr_items: list[dict]) -> dict:
    prev_tasks = {it.get('task') for it in prev_items}
    prev_by_key: dict = {}
    for it in prev_items:
        prev_by_key.setdefault((it.get('task'), it.get('countryCode')), it)
    changes: dict = {}
    summary: dict = {}
    seen: set = set()
    for it in curr_items:
        key = (it.get('task'), it.get('countryCode'))
        if key in seen:
            continue
        seen.add(key)
        change = None
        if it.get('task') not in prev_tasks:
            change = {'type': 'new_task'}
        elif key not in prev_by_key:
            change = {'type': 'new_country'}
        else:
            p = prev_by_key[key]
            if p.get('status') != it.get('status'):
                change = {'type': 'status', 'from': p.get('status'), 'to': it.get('status')}
            elif p.get('url') != it.get('url'):
                change = {'type': 'url'}
        if change is None:
            continue
        changes[f'{key[0]}|{key[1]}'] = change
        s = summary.setdefault(key[0], {'changed': 0, 'new': 0})
        s['new' if change['type'] in ('new_task', 'new_country') else 'changed'] += 1
    return {'changes': changes, 'taskSummary': summary}


def gr(prev_path: Path, prev_date: str) -> int:
    titles = json.loads(GR_TITLES.read_text(encoding='utf-8')).get('titles') or {}
    t2t = {v: k for k, v in titles.items()}
    try:
        curr = read_gr_items(GR_XLSX, t2t)
    except Exception as e:  # noqa: BLE001
        _summary(f'## ❌ Global Request 업로드 실패\n\n엑셀을 열 수 없습니다: `{e}`\n\n이전 파일로 되돌렸습니다.')
        return 1
    if not curr:
        _summary("## ❌ Global Request 업로드 실패\n\n'Head' 머리글 행을 가진 시트가 하나도 없습니다. 양식이 맞는지 확인해 주세요.\n\n이전 파일로 되돌렸습니다.")
        return 1
    prev = read_gr_items(prev_path, t2t) if prev_path.exists() and prev_path.stat().st_size else []
    cmap = normalize_countries([i['countryRaw'] for i in prev + curr])
    for it in prev + curr:
        it['countryCode'] = cmap.get(it['countryRaw'], '')
    diff = diff_gr_items(prev, curr)
    payload = {'baseDate': datetime.now(KST).strftime('%Y%m%d'), 'prevDate': prev_date,
               'changes': diff['changes'], 'taskSummary': diff['taskSummary'],
               'generatedAt': datetime.now(KST).date().isoformat()}
    GR_CHANGES.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    unresolved = sorted({i['countryRaw'] for i in curr if not i['countryCode']})
    lines = [f'## ✅ Global Request 반영', '',
             f'- 태스크 {len({i["task"] for i in curr})}개 · 국가행 {len(curr)}건',
             f'- 금주 변경 배지: {len(diff["changes"])}건 (기준 {prev_date} → 오늘)']
    for t, s in sorted(diff['taskSummary'].items(), key=lambda x: -(x[1]['new'] + x[1]['changed'])):
        lines.append(f'  - {t}: 변경 {s["changed"]} · 신규 {s["new"]}')
    if unresolved:
        lines.append(f'- ⚠️ 국가명을 코드로 바꾸지 못함 {len(unresolved)}건 (배지가 안 붙음): {", ".join(unresolved[:10])}')
    _summary('\n'.join(lines))
    return 0


def main() -> int:
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest='cmd', required=True)
    sub.add_parser('export')
    sub.add_parser('import')
    g = sub.add_parser('gr')
    g.add_argument('prev', type=Path)
    g.add_argument('--prev-date', required=True)
    a = ap.parse_args()
    if a.cmd == 'export':
        export_status()
        return 0
    if a.cmd == 'import':
        return import_status()
    return gr(a.prev, a.prev_date)


if __name__ == '__main__':
    sys.exit(main())
